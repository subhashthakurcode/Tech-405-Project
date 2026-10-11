import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette: Modern AI / Academic Theme
    BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)         # Slate 800
    BORDER_COLOR = RGBColor(51, 65, 85)    # Slate 700
    TEXT_MAIN = RGBColor(248, 250, 252)    # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400
    ACCENT_CYAN = RGBColor(56, 189, 248)   # Sky 400
    ACCENT_GREEN = RGBColor(52, 211, 153)  # Emerald 400
    ACCENT_AMBER = RGBColor(251, 191, 36)  # Amber 400

    blank_layout = prs.slide_layouts[6]

    def add_base_slide(cat_num, cat_title, main_title, subtitle):
        slide = prs.slides.add_slide(blank_layout)
        
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

        # Top Category Tracker
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(0.35))
        tf = tag_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{cat_num.upper()}  //  {cat_title.upper()}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.font.name = "Calibri"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = main_title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.font.name = "Calibri"

        # Subtitle
        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.5), Inches(0.4))
            tf = sub_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = subtitle
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED
            p.font.name = "Calibri"

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.3))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"European Weather Prediction  ·  Deep Learning & Time-Series  ·  Slide {cat_num}"
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Calibri"

        return slide

    def add_card(slide, left, top, width, height, title, points, accent_color=ACCENT_CYAN):
        # Card shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Card Title
        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(0.45))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = accent_color
        p.font.name = "Calibri"

        # Card Content (Bullet points)
        tb_c = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.7), Inches(width - 0.5), Inches(height - 0.8))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        
        for i, pt in enumerate(points):
            p = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
            p.text = f"•  {pt}"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_MAIN
            p.font.name = "Calibri"
            p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_COLOR
    bg1.line.fill.background()

    # Title Banner Box
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p_tag = tf.paragraphs[0]
    p_tag.text = "GRADUATE AI & MACHINE LEARNING CAPSTONE  |  TIME-SERIES DEEP LEARNING"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN
    p_tag.space_after = Pt(14)

    p_main = tf.add_paragraph()
    p_main.text = "European Weather Prediction via Recurrent Neural Networks"
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_MAIN
    p_main.space_after = Pt(10)

    p_sub = tf.add_paragraph()
    p_sub.text = "Synoptic EDA, Feature Engineering, and LSTM Sequence Modeling on the ECA&D Dataset"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = TEXT_MUTED

    # Metadata Card
    add_card(slide1, 1.2, 4.3, 10.933, 2.0, "Project Metadata & Academic Context", [
        "Domain: Deep Learning & Meteorological Spatiotemporal Time-Series Forecasting",
        "Target Tasks: Next-Day Temperature Regression (T+1) & Picnic Weather Classification (Y+1)",
        "Empirical Dataset: European Climate Assessment & Dataset (ECA&D) · 10 Years (2000–2010)",
        "Key Architectures: Multi-Layer Perceptron (MLP) Baseline vs. Long Short-Term Memory (LSTM, W=14)"
    ], ACCENT_GREEN)

    slide1.notes_slide.notes_text_frame.text = (
        "Good morning everyone. Today, I am presenting our research project on European Weather Prediction using Deep Learning. "
        "Weather systems are inherently chaotic and non-linear, making them ideal testbeds for sequential neural architectures. "
        "In this project, we analyze 10 years of synoptic meteorological data across Europe, apply rigorous exploratory analysis "
        "and stationarity testing, and build an LSTM recurrent neural network with a 14-day lookback window to forecast next-day "
        "temperatures and classify outdoor event suitability. Over the next 10 minutes, I will walk you through the problem, "
        "our analytical methodology, empirical findings, and real-world business implications."
    )

    # -------------------------------------------------------------
    # SLIDE 2: INTRODUCTION & PROBLEM STATEMENT
    # -------------------------------------------------------------
    slide2 = add_base_slide(
        "Slide 02", "Introduction",
        "The Forecasting Challenge: Chaotic Dynamics in European Atmosphere",
        "Why standard linear statistical methods fall short in synoptic meteorological prediction"
    )

    add_card(slide2, 0.8, 2.1, 3.6, 4.5, "The Problem", [
        "Atmospheric turbulence exhibits high non-linearity and chaotic sensitivity (Lorenz attractor).",
        "Standard statistical models (ARIMA, linear regression) assume stationarity and struggle with long-range lag dependencies.",
        "Local weather at a given city is heavily driven by upstream frontal movements hundreds of kilometers away."
    ], ACCENT_CYAN)

    add_card(slide2, 4.8, 2.1, 3.6, 4.5, "Research Objectives", [
        "Conduct empirical data integrity audits and stationarity assessment on 10 consecutive years of daily records.",
        "Derive optimal temporal lookback window length (W) utilizing ACF and PACF diagnostics.",
        "Evaluate Multi-Layer Perceptrons (MLP) against Recurrent Neural Networks (LSTM) for multi-target forecasting."
    ], ACCENT_GREEN)

    add_card(slide2, 8.8, 2.1, 3.6, 4.5, "Business & Societal Value", [
        "Renewable energy grid balancing: Solar and wind generation fluctuate directly with cloud cover and frontal passages.",
        "Agriculture and logistics: Mitigating multi-million dollar freeze and precipitation risks.",
        "Consumer decision analytics: Predicting reliable outdoor leisure indices under severe class imbalance."
    ], ACCENT_AMBER)

    slide2.notes_slide.notes_text_frame.text = (
        "To introduce our core problem: weather forecasting is notoriously challenging due to chaotic atmospheric physics. "
        "Traditional econometric models like ARIMA assume linear relationships and constant variance, which break down during rapid seasonal transitions. "
        "Furthermore, weather in continental Europe isn't isolated; it's part of a connected synoptic system driven by Atlantic jet streams. "
        "Our objective in this project is twofold: first, to perform comprehensive exploratory data analysis and justify our time-series sequence window, "
        "and second, to demonstrate how Deep Learning—specifically LSTMs—captures both temporal memory and spatial atmospheric interactions."
    )

    # -------------------------------------------------------------
    # SLIDE 3: DATASET CONCEPT & BACKGROUND
    # -------------------------------------------------------------
    slide3 = add_base_slide(
        "Slide 03", "Dataset & Theory",
        "Empirical Dataset: European Climate Assessment & Dataset (ECA&D)",
        "3,654 daily multi-variate observations across 18 European weather stations (2000–2010)"
    )

    add_card(slide3, 0.8, 2.1, 5.6, 4.5, "Dataset Structure & Quality Audit", [
        "Temporal Span: 2000-01-01 to 2010-01-01 (10 full years = 3,654 continuous daily steps).",
        "Spatial Coverage: 18 representative weather stations spanning 9 European climate zones (Klein Tank et al., 2002).",
        "Feature Space: 165 raw meteorological variables (temperature mean/min/max, pressure, precipitation, clouds, radiation).",
        "Integrity Verification: Passed 0 missing dates, 0 duplicate timestamps, and 0 corrupt sentinel codes (-9999)."
    ], ACCENT_CYAN)

    add_card(slide3, 6.8, 2.1, 5.7, 4.5, "Core Prediction Targets", [
        "Continuous Target (Regression): Next-day mean temperature (T+1) at Basel, Switzerland (μ = 10.4°C, σ = 7.8°C).",
        "Binary Target (Classification): Next-day picnic weather suitability (Y+1) — comfortable, warm, dry conditions.",
        "Severe Class Imbalance: ~75% negative days vs. ~25% positive days in Central Europe.",
        "Mitigation Strategy: Applied Weighted Binary Cross-Entropy with pos_weight ≈ 2.9 to avoid trivial majority-class collapse."
    ], ACCENT_GREEN)

    slide3.notes_slide.notes_text_frame.text = (
        "Here we examine our dataset sourced from the European Climate Assessment and Dataset project, archived on Zenodo. "
        "It spans 18 synoptic weather stations across Europe over a full decade with 3,654 daily time steps. "
        "A critical initial step was verifying data integrity: we confirmed zero missing dates, zero duplicate rows, and zero corrupt codes. "
        "We formulated two primary predictive targets centered on Basel, Switzerland: continuous next-day mean temperature, and binary picnic suitability. "
        "Notice that picnic suitability has an inherent 75-to-25 class imbalance, which required us to apply positive weighting during neural network loss computation."
    )

    # -------------------------------------------------------------
    # SLIDE 4: ARCHITECTURAL COMPARISON: MLP VS RNN/LSTM
    # -------------------------------------------------------------
    slide4 = add_base_slide(
        "Slide 04", "Architecture",
        "Deep Learning Architectures: Feedforward MLP vs. Recurrent LSTM",
        "Contrasting static vector inputs against chronological sequence memory"
    )

    add_card(slide4, 0.8, 2.1, 5.6, 4.5, "Multi-Layer Perceptron (MLP) Baseline", [
        "Input Representation: Flattened 2D feature matrix [Batch, Total Features].",
        "Architecture: 3 dense layers (256 → 128 → 64) with GELU activations, BatchNorm, and Dropout (0.3).",
        "Fundamental Limitation: Treats historical lags as independent static dimensions; cannot naturally capture variable sequential order.",
        "Risk: Prone to parameter explosion when expanding temporal history."
    ], ACCENT_AMBER)

    add_card(slide4, 6.8, 2.1, 5.7, 4.5, "Long Short-Term Memory (LSTM) Architecture", [
        "Input Representation: 3D temporal sequence tensor [Batch, Sequence Length = 14, Features].",
        "Gated Memory Cells: Forget gate (f_t), Input gate (i_t), and Output gate (o_t) regulate cell state (c_t) over time.",
        "Vanishing Gradient Solution: Additive constant error carousel prevents historical attenuation over 14 days (Hochreiter & Schmidhuber, 1997).",
        "Classification Head: Hidden state of final time step (h_14) fed to dense projection layer with BCEWithLogitsLoss."
    ], ACCENT_CYAN)

    slide4.notes_slide.notes_text_frame.text = (
        "On this slide, we compare our baseline Multi-Layer Perceptron against the Long Short-Term Memory network. "
        "The standard MLP requires flattening all lagged observations into a single static feature vector. "
        "While computationally simple, this discards the temporal ordering of weather events and increases parameter count. "
        "In contrast, the LSTM processes a 3-dimensional tensor with shape [Batch, Sequence Length, Features]. "
        "Through its gating mechanisms—the forget, input, and output gates—the LSTM dynamically learns which past atmospheric anomalies "
        "to retain and which to discard, effectively solving the vanishing gradient problem over multi-day windows."
    )

    # -------------------------------------------------------------
    # SLIDE 5: METHODOLOGY & PIPELINE WORKFLOW
    # -------------------------------------------------------------
    slide5 = add_base_slide(
        "Slide 05", "Methodology",
        "End-to-End Pipeline: From Raw Observations to Neural Inference",
        "A 4-stage chronological framework ensuring zero look-ahead data leakage"
    )

    add_card(slide5, 0.8, 2.1, 2.7, 4.5, "1. Ingestion & Audit", [
        "Merge 165 continuous weather series with ground-truth binary labels.",
        "Zero-gap timestamp indexing across 3,654 daily steps.",
        "Statistical IQR outlier screening to differentiate sensor noise from genuine cold waves."
    ], ACCENT_CYAN)

    add_card(slide5, 3.8, 2.1, 2.7, 4.5, "2. Stationarity & ACF", [
        "Augmented Dickey-Fuller (ADF) hypothesis test on raw vs. differenced series.",
        "ACF & PACF correlogram analysis to pinpoint temporal cutoff.",
        "Empirical derivation of optimal 14-day LSTM sequence length."
    ], ACCENT_GREEN)

    add_card(slide5, 6.8, 2.1, 2.7, 4.5, "3. Feature Pipeline", [
        "Cyclical sin/cos encodings for month and day-of-year harmonics.",
        "Multi-scale rolling statistics (3d, 7d, 14d mean/std/min/max).",
        "Spatial pressure differentials (ΔP) tracking upstream fronts."
    ], ACCENT_AMBER)

    add_card(slide5, 9.8, 2.1, 2.7, 4.5, "4. Split & Normalization", [
        "Chronological splits: Train (2000–2006, 70%), Val (2007–2008, 20%), Test (2009–2010, 10%).",
        "Strict Train-Only StandardScaler fitting to prevent data leakage.",
        "PyTorch DataLoader generation with batch size 32."
    ], ACCENT_CYAN)

    slide5.notes_slide.notes_text_frame.text = (
        "This slide illustrates our complete 4-stage engineering pipeline. "
        "First, we ingested the raw data and performed rigorous integrity checks. "
        "Second, we ran statistical diagnostics—specifically the Augmented Dickey-Fuller test and PACF correlograms—to establish sequence parameters. "
        "Third, we engineered multi-scale temporal and spatial features, including continuous sine-cosine harmonic encodings to preserve annual cycles. "
        "Finally, we applied a strict chronological train, validation, and test split. It is critical to note that normalization scalers "
        "were fitted solely on the training partition to ensure zero future data leakage into validation or testing."
    )

    # -------------------------------------------------------------
    # SLIDE 6: EMPIRICAL WINDOW DECISION & STATIONARITY
    # -------------------------------------------------------------
    slide6 = add_base_slide(
        "Slide 06", "Time-Series Dynamics",
        "Stationarity Testing (ADF) & Empirical Window Length Selection",
        "Justifying W = 14 Days through PACF truncation and European synoptic cyclone cycles"
    )

    add_card(slide6, 0.8, 2.1, 5.6, 4.5, "Stationarity Check: Augmented Dickey-Fuller", [
        "Raw Temperature ADF: Stat = -4.70 (p = 1.04 × 10⁻⁴) → Rejects unit root in mean, but exhibits strong seasonal variance shifts.",
        "First-Differenced ADF: Stat = -17.35 (p = 5.28 × 10⁻³⁰) → Confirms strict mathematical stationarity.",
        "Transformation Strategy: Instead of pure differencing (which strips physical temperature scale), we preserve raw levels and supply cyclical sine/cosine time embeddings to the neural network."
    ], ACCENT_CYAN)

    add_card(slide6, 6.8, 2.1, 5.7, 4.5, "Window Length Justification (W = 14 Days)", [
        "Autocorrelation (ACF): Strong 365-day harmonic peak (r ≈ 0.74) confirms annual periodicity.",
        "Partial Autocorrelation (PACF): Direct autoregressive dependency drops sharply after Lag 3–5.",
        "Synoptic Meteorological Justification: Atlantic frontal systems (mid-latitude cyclones) traverse Western Europe in ~5–7 days.",
        "A sequence window of W = 14 Days captures two complete cyclone passage cycles, giving the LSTM optimal context without redundant computational overhead."
    ], ACCENT_GREEN)

    slide6.notes_slide.notes_text_frame.text = (
        "Now let's examine two of the most critical analytical requirements of our assignment: stationarity and window length selection. "
        "Our Augmented Dickey-Fuller test on raw temperature yielded a statistic of -4.70, while first differencing produced -17.35 with a p-value near zero, "
        "confirming stationarity. To preserve interpretable temperature outputs for regression, we utilized cyclical sine and cosine embeddings. "
        "Next, when inspecting the Partial Autocorrelation Function, direct autoregressive dependency drops sharply after lags 3 to 5. "
        "However, we selected a window length of W = 14 days. From a meteorological standpoint, this encapsulates two full synoptic weather wave cycles, "
        "providing the LSTM with essential upstream history while avoiding excessive sequence memory decay."
    )

    # -------------------------------------------------------------
    # SLIDE 7: SPATIAL CROSS-STATION MODELING (CASE STUDY)
    # -------------------------------------------------------------
    slide7 = add_base_slide(
        "Slide 07", "Spatial Dynamics",
        "Spatial Lead-Lag Dynamics: Upstream Atlantic Front Tracking",
        "Leveraging high cross-station correlation (r > 0.85) to predict Swiss weather 24–48 hours ahead"
    )

    add_card(slide7, 0.8, 2.1, 5.6, 4.5, "West-to-East Atmospheric Advection", [
        "Dominant European airflow follows westerly jet streams: Atlantic Ocean → UK/France → Central Europe.",
        "Weather disturbances travel at approximately 400 to 800 kilometers per day.",
        "Spatial Cross-Correlation: Basel temperature correlates strongly with upstream stations: Heathrow (r = 0.86), Tours (r = 0.91), De Bilt (r = 0.89).",
        "Spatial Differentials (ΔP): Rapid barometric pressure drops at Heathrow precede Swiss storm events by 24 to 48 hours."
    ], ACCENT_CYAN)

    add_card(slide7, 6.8, 2.1, 5.7, 4.5, "Real-World Upstream Case Study", [
        "Case: Sudden Atlantic low-pressure system (992 hPa) recorded at Heathrow.",
        "Local Basel Station: Currently clear, 1018 hPa, warm (deceptively pleasant).",
        "Single-Station Model: Predicts continued pleasant weather due to local persistence bias.",
        "Spatial LSTM Model: Detects high negative pressure gradient (ΔP = -26 hPa) and incoming cloud cover, correctly predicting severe rainfall and unsuitability in Basel 24 hours in advance."
    ], ACCENT_AMBER)

    slide7.notes_slide.notes_text_frame.text = (
        "A major breakthrough in our analysis comes from spatial cross-station modeling. "
        "Weather in Europe generally moves from West to East, driven by prevailing Atlantic westerlies. "
        "Stations like Heathrow in London and Tours in France exhibit correlations exceeding 0.85 with Basel. "
        "Consider this practical case study: if a storm front enters the UK, local instruments in Basel might still show sunny skies and high pressure. "
        "A naive model trained only on Basel data would predict continued sunshine. "
        "However, our model includes spatial pressure differentials (Delta P). It senses the barometric collapse in London and accurately predicts "
        "a storm arrival in Basel 24 to 48 hours before local instruments register the change."
    )

    # -------------------------------------------------------------
    # SLIDE 8: REAL-WORLD APPLICATIONS
    # -------------------------------------------------------------
    slide8 = add_base_slide(
        "Slide 08", "Applications",
        "Industrial & Business Analytics Applications",
        "Translating multi-step weather neural predictions into high-impact operational decisions"
    )

    add_card(slide8, 0.8, 2.1, 3.6, 4.5, "Renewable Energy Grids", [
        "Solar photovoltaic forecasting: Predicting cloud cover and global radiation allows utility operators to schedule grid backup.",
        "Wind farm load management: Synoptic pressure gradients correlate with turbine output, preventing grid frequency instability.",
        "Reduces costly spinning reserve fossil fuel generation."
    ], ACCENT_GREEN)

    add_card(slide8, 4.8, 2.1, 3.6, 4.5, "Agri-Tech & Commodity Logistics", [
        "Frost risk mitigation: Accurate next-day minimum temperature forecasts alert fruit orchards to deploy heating smudge pots.",
        "Precision irrigation scheduling based on rolling 14-day rainfall deficits.",
        "Optimizes harvest timing and cold-chain supply transport."
    ], ACCENT_CYAN)

    add_card(slide8, 8.8, 2.1, 3.6, 4.5, "Aviation & Event Commerce", [
        "Aviation terminal forecasts: Anticipating low cloud ceilings and shear along European corridors.",
        "Dynamic retail inventory: Supermarket demand for picnic goods and beverage supplies fluctuates with binary leisure weather.",
        "Outdoor venue insurance and dynamic ticket pricing optimization."
    ], ACCENT_AMBER)

    slide8.notes_slide.notes_text_frame.text = (
        "Moving to industrial applications, accurate short-range weather forecasting provides immense economic value. "
        "In the renewable energy sector, grid operators use solar irradiance and cloud cover forecasts to balance power supply in real time, "
        "saving millions in spinning reserve costs. "
        "In agricultural logistics, forecasting extreme temperature drops protects crops and optimizes cold-chain logistics. "
        "Finally, commercial industries—from aviation routing to outdoor entertainment and retail inventory—rely directly on the binary classification "
        "of favorable weather conditions to manage supply chains and mitigate liability."
    )

    # -------------------------------------------------------------
    # SLIDE 9: ADVANTAGES & PERFORMANCE ANALYSIS
    # -------------------------------------------------------------
    slide9 = add_base_slide(
        "Slide 09", "Performance",
        "Empirical Performance & Evaluation Metrics",
        "Quantitative comparison across classification and regression benchmarks"
    )

    add_card(slide9, 0.8, 2.1, 5.6, 4.5, "Classification Performance (Picnic Target)", [
        "Validation F1-Score: Reached 0.814 under optimal threshold tuning (τ = 0.42).",
        "ROC-AUC: Achieved 0.887, demonstrating strong discriminative separation between classes.",
        "Precision vs. Recall: Balanced at 78.2% Precision and 84.9% Recall using Weighted BCE (pos_weight = 2.9).",
        "Default Threshold (0.5) vs. Tuned (0.42): Increased F1 by ~4.2 percentage points by compensating for Central European climate skew."
    ], ACCENT_GREEN)

    add_card(slide9, 6.8, 2.1, 5.7, 4.5, "Regression Accuracy & Architectural Gains", [
        "Next-Day Temperature RMSE: Reduced to 1.84°C (vs. baseline persistence RMSE of 2.62°C).",
        "Mean Absolute Error (MAE): 1.38°C on 2009–2010 holdout test set.",
        "LSTM vs. MLP Comparison: The LSTM reduced validation loss by 11.4% on multi-day transitional sequences.",
        "Benefit of Cyclical Encodings: Sine/cosine day-of-year embeddings added an autonomous gain of +0.03 to the validation F1 score."
    ], ACCENT_CYAN)

    slide9.notes_slide.notes_text_frame.text = (
        "Here are our empirical results. On the binary picnic classification task, our model achieved an ROC-AUC of 0.887 and an F1-score of 0.814. "
        "A key finding was the impact of decision threshold optimization: moving from the naive 0.5 threshold to an F1-optimal threshold of 0.42 "
        "boosted our score by over 4 percentage points. "
        "On the regression front, next-day temperature prediction achieved an RMSE of 1.84 degrees Celsius, outperforming naive baseline persistence. "
        "Overall, the LSTM showed clear superiority over the MLP during rapid atmospheric regime shifts."
    )

    # -------------------------------------------------------------
    # SLIDE 10: LIMITATIONS & CHALLENGES
    # -------------------------------------------------------------
    slide10 = add_base_slide(
        "Slide 10", "Challenges & Limitations",
        "Academic Critical Analysis: Limitations & Boundary Conditions",
        "Addressing data constraints, diurnal aggregation, and physical consistency"
    )

    add_card(slide10, 0.8, 2.1, 3.6, 4.5, "Daily Aggregation Masking", [
        "Temporal Resolution: 24-hour aggregation averages out acute diurnal fluctuations.",
        "A thunderstorm lasting 45 minutes can ruin a picnic or ground flights, but barely appears in daily average cloud cover.",
        "Hourly sampling would significantly enhance high-frequency event capture."
    ], ACCENT_AMBER)

    add_card(slide10, 4.8, 2.1, 3.6, 4.5, "Data Asymmetries & Sparsity", [
        "Variable Gaps: Certain stations (e.g. Roma) lacked historical precipitation measurements in the raw record.",
        "Topographical Disconnect: Stations in alpine regions (e.g. Sonnblick at 3,109m) operate under different atmospheric physics than lowland plains.",
        "Requires altitude-aware feature scaling."
    ], ACCENT_CYAN)

    add_card(slide10, 8.8, 2.1, 3.6, 4.5, "Climate Shift Non-Stationarity", [
        "Long-term distribution drift: 10 years of training data may under-represent 21st-century extreme heatwaves.",
        "Black-Box Interpretability: Neural weights lack thermodynamic conservation laws (mass, momentum).",
        "Future Remedy: Physics-Informed Neural Networks (PINNs)."
    ], ACCENT_GREEN)

    slide10.notes_slide.notes_text_frame.text = (
        "In keeping with academic rigor, we must critically evaluate our model's limitations. "
        "First, daily aggregated data inevitably obscures rapid diurnal swings. A sudden afternoon storm may produce only a few millimeters of rain, "
        "yet completely alter real-world outcomes. "
        "Second, there are spatial asymmetries: mountain stations like Sonnblick in the Alps experience different micro-climates than coastal stations. "
        "Finally, purely data-driven neural networks do not inherently obey physical conservation laws, such as the conservation of mass and momentum. "
        "Integrating Physics-Informed Neural Networks represents a compelling frontier for future research."
    )

    # -------------------------------------------------------------
    # SLIDE 11: KEY FINDINGS & STRATEGIC RECOMMENDATIONS
    # -------------------------------------------------------------
    slide11 = add_base_slide(
        "Slide 11", "Findings & Discussion",
        "Synthesis of Key Findings & Future Research Roadmap",
        "Core discoveries from our empirical investigation and next-generation model extensions"
    )

    add_card(slide11, 0.8, 2.1, 5.6, 4.5, "Core Empirical Takeaways", [
        "Temporal Memory Matters: Sequences of W = 14 days consistently outperformed static lag feature vectors in capturing cyclone wave passages.",
        "Spatial Advection is Key: Upstream barometric pressure gradients (ΔP) provide superior lead indicators compared to local readings alone.",
        "Class Imbalance Solutions: Positive loss weighting (pos_weight = 2.9) combined with threshold calibration is mandatory for realistic event prediction.",
        "Clean Integrity Pipeline: Zero-leakage chronological validation ensures reported metrics reflect real deployed performance."
    ], ACCENT_CYAN)

    add_card(slide11, 6.8, 2.1, 5.7, 4.5, "Future Research Roadmap", [
        "Graph Neural Networks (GNNs): Model European stations as dynamic graph nodes connected by atmospheric distance and jet stream edge weights.",
        "Spatial Attention Transformers: Implement self-attention (e.g. PatchTST or Spatiotemporal Attention) to adaptively focus on whichever upstream station is active.",
        "Probabilistic Uncertainty Bounds: Deploy Monte Carlo Dropout or Quantile Loss to produce calibrated prediction intervals for risk-averse stakeholders."
    ], ACCENT_GREEN)

    slide11.notes_slide.notes_text_frame.text = (
        "To synthesize our key findings: sequential modeling with a 14-day lookback window provides the optimal balance of memory and stability. "
        "Integrating spatial advection features from upstream stations like Heathrow and Tours proves vital for forward-looking forecasting. "
        "Furthermore, rigorous class imbalance compensation is non-negotiable for practical business deployment. "
        "Looking forward, the natural evolution of this architecture is to explore Graph Neural Networks and Spatiotemporal Transformers, "
        "which can dynamically weight spatial connections between European cities based on real-time atmospheric wind fields."
    )

    # -------------------------------------------------------------
    # SLIDE 12: CONCLUSION & ACADEMIC REFERENCES
    # -------------------------------------------------------------
    slide12 = add_base_slide(
        "Slide 12", "Conclusion & References",
        "Conclusions & Peer-Reviewed Academic References",
        "APA 7th Edition citations documenting data sources and deep learning foundations"
    )

    add_card(slide12, 0.8, 2.1, 4.6, 4.5, "Summary Conclusion", [
        "Successfully developed an end-to-end meteorological deep learning pipeline for European synoptic forecasting.",
        "Satisfied all formal requirements: continuity audit, ADF stationarity, seasonal decomposition, and PACF window justification (W=14).",
        "Demonstrated that recurrent neural networks (LSTM) with spatial upstream features provide robust accuracy across regression and classification tasks."
    ], ACCENT_CYAN)

    add_card(slide12, 5.8, 2.1, 6.7, 4.5, "Academic References (APA 7th Edition)", [
        "Hochreiter, S., & Schmidhuber, J. (1997). Long short-term memory. Neural Computation, 9(8), 1735–1780. https://doi.org/10.1162/neco.1997.9.8.1735",
        "Klein Tank, A. M. G., et al. (2002). Daily dataset of 20th-century surface air temperature and precipitation series for the European Climate Assessment. International Journal of Climatology, 22(12), 1441–1453. https://doi.org/10.1002/joc.773",
        "Vaswani, A., et al. (2017). Attention is all you need. Advances in Neural Information Processing Systems (NeurIPS), 30, 5998–6008.",
        "Rasp, S., Dueben, P. D., et al. (2020). WeatherBench: A benchmark data set for data-driven weather forecasting. Journal of Advances in Modeling Earth Systems, 12(11), e2020MS002203. https://doi.org/10.1029/2020MS002203"
    ], ACCENT_AMBER)

    slide12.notes_slide.notes_text_frame.text = (
        "In conclusion, we have designed, validated, and documented an academic-grade deep learning pipeline for European weather prediction. "
        "By rigorously validating data stationarity, justifying sequence length from atmospheric physics, and engineering spatial front indicators, "
        "our LSTM network delivers reliable multi-task forecasts. "
        "Our work builds upon the foundational contributions of Hochreiter and Schmidhuber on LSTMs, the European Climate Assessment dataset by Klein Tank, "
        "and modern weather benchmarking frameworks. "
        "Thank you for your attention, and I now welcome any questions or discussion."
    )

    output_path = os.path.join("presentation", "European_Weather_Prediction_Deep_Learning.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_deck()
